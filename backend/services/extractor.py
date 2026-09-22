"""
Text extraction service for PDF and DOCX resume documents.
Extracts raw text purely in-memory without saving files permanently to disk.
"""
import io
from pypdf import PdfReader
from docx import Document


class ExtractionError(Exception):
    """Custom exception raised when resume text extraction fails."""
    pass


def extract_text_from_pdf(file_bytes: bytes) -> str:
    """
    Extracts text from PDF byte stream using pypdf.
    """
    try:
        pdf_file = io.BytesIO(file_bytes)
        reader = PdfReader(pdf_file)
        
        if len(reader.pages) == 0:
            raise ExtractionError("The PDF document does not contain any pages.")

        extracted_text = []
        for page_idx, page in enumerate(reader.pages):
            text = page.extract_text()
            if text:
                extracted_text.append(text)

        full_text = "\n\n".join(extracted_text).strip()

        if not full_text:
            raise ExtractionError(
                "No readable text found in the PDF. If this is a scanned image or photo of a resume, "
                "please provide a PDF with selectable text or a DOCX document."
            )

        return full_text
    except ExtractionError:
        raise
    except Exception as e:
        raise ExtractionError(f"Failed to parse PDF document: {str(e)}")


def extract_text_from_docx(file_bytes: bytes) -> str:
    """
    Extracts text from DOCX byte stream using python-docx.
    """
    try:
        docx_file = io.BytesIO(file_bytes)
        doc = Document(docx_file)

        extracted_paragraphs = []
        # Extract normal paragraphs
        for p in doc.paragraphs:
            clean_p = p.text.strip()
            if clean_p:
                extracted_paragraphs.append(clean_p)

        # Extract text from tables if any
        for table in doc.tables:
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_text:
                    extracted_paragraphs.append(" | ".join(row_text))

        full_text = "\n".join(extracted_paragraphs).strip()

        if not full_text:
            raise ExtractionError(
                "The Word document (.docx) does not contain any readable text."
            )

        return full_text
    except ExtractionError:
        raise
    except Exception as e:
        raise ExtractionError(f"Failed to parse Word (.docx) document: {str(e)}")


def extract_resume_text(file_storage, extension: str) -> str:
    """
    Dispatcher to extract text based on file extension.
    Reads file stream into bytes and parses in-memory.
    """
    file_bytes = file_storage.read()
    file_storage.seek(0)  # Reset stream

    ext = extension.lower()
    if ext == '.pdf':
        return extract_text_from_pdf(file_bytes)
    elif ext == '.docx':
        return extract_text_from_docx(file_bytes)
    else:
        raise ExtractionError(f"Unsupported file extension: {extension}")
