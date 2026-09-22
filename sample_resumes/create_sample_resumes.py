"""
Utility script to generate sample PDF and DOCX resumes for testing.
"""
import os
import io
from docx import Document
from docx.shared import Inches, Pt, RGBColor

SAMPLE_DIR = os.path.dirname(os.path.abspath(__file__))

def create_sample_docx():
    doc = Document()

    # Title / Name
    title_p = doc.add_paragraph()
    title_run = title_p.add_run("Alex Morgan")
    title_run.font.size = Pt(20)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(75, 0, 130)

    # Contact Info
    contact_p = doc.add_paragraph("alex.morgan@email.com | +1 (555) 234-5678 | San Francisco, CA | linkedin.com/in/alexmorgan | github.com/alexmorgan")
    contact_p.runs[0].font.size = Pt(9.5)
    contact_p.runs[0].font.italic = True

    # Summary
    doc.add_heading("Professional Summary", level=1)
    doc.add_paragraph(
        "Passionate Full Stack Software Engineer with 2+ years of hands-on experience designing and scaling web applications, "
        "REST APIs, and machine learning microservices. Proficient in Python, React, Flask, and PostgreSQL. "
        "Adept at agile collaboration, performance tuning, and building clean, accessible user interfaces."
    )

    # Technical Skills
    doc.add_heading("Technical Skills", level=1)
    doc.add_paragraph(
        "Languages: Python, JavaScript (ES6+), TypeScript, SQL, HTML5, CSS3\n"
        "Frameworks & Libraries: React, Node.js, Flask, Tailwind CSS, Express, Next.js\n"
        "Databases & Cloud: PostgreSQL, MongoDB, Redis, Docker, AWS (S3, EC2), Git\n"
        "Core Competencies: RESTful API Design, Microservices, Agile Scrum, System Architecture"
    )

    # Experience
    doc.add_heading("Work Experience", level=1)
    
    p1 = doc.add_paragraph()
    r1 = p1.add_run("Software Engineer Intern | NovaCloud Solutions")
    r1.bold = True
    p1.add_run(" (June 2023 – Present)\n")
    p1.add_run("• Developed and deployed 6+ RESTful microservice endpoints using Python and Flask, reducing client latency by 28%.\n")
    p1.add_run("• Implemented responsive frontend dashboards in React and Tailwind CSS, increasing daily active user engagement by 35%.\n")
    p1.add_run("• Automated CI/CD deployment pipelines using GitHub Actions and Docker, cutting staging deployment cycles in half.")

    p2 = doc.add_paragraph()
    r2 = p2.add_run("Frontend Developer Co-op | PulseTech Interactive")
    r2.bold = True
    p2.add_run(" (Jan 2023 – May 2023)\n")
    p2.add_run("• Built modular reusable UI components for an internal data visualization portal using React and Chart.js.\n")
    p2.add_run("• Collaborated with UX designers to achieve WCAG 2.1 AA accessibility standards across 15+ pages.")

    # Projects
    doc.add_heading("Projects", level=1)
    
    doc.add_paragraph(
        "AI Content Assistant (React, Python, Flask, Gemini API)\n"
        "• Engineered a generative writing platform that drafts blog posts and summarizes documents.\n"
        "• Implemented asynchronous streaming responses and in-memory caching to support 500+ daily requests."
    )

    doc.add_paragraph(
        "Real-Time Collaborative Code Editor (React, WebSocket, Node.js)\n"
        "• Built an interactive in-browser coding playground featuring real-time collaborative editing and syntax highlighting.\n"
        "• Integrated operational transformation algorithms for seamless concurrent typing conflict resolution."
    )

    # Education
    doc.add_heading("Education", level=1)
    doc.add_paragraph(
        "Bachelor of Science in Computer Science | State University of Technology\n"
        "Graduation: May 2024 | GPA: 3.82 / 4.0\n"
        "Coursework: Data Structures, Algorithms, Database Systems, Web Development, Computer Networks"
    )

    # Certifications
    doc.add_heading("Certifications", level=1)
    doc.add_paragraph(
        "• AWS Certified Cloud Practitioner (2024)\n"
        "• Meta Certified Front-End Developer (2023)"
    )

    output_path = os.path.join(SAMPLE_DIR, "sample_resume.docx")
    doc.save(output_path)
    print(f"Created sample DOCX resume at: {output_path}")
    return output_path


def create_sample_pdf():
    """
    Creates a valid standard PDF document without external C libraries.
    Constructs a clean PDF 1.4 stream containing resume text.
    """
    text_content = """BT
/F1 18 Tf
50 780 Td
(Alex Morgan - Full Stack Engineer) Tj
/F1 10 Tf
0 -18 Td
(alex.morgan@email.com | +1 (555) 234-5678 | San Francisco, CA | github.com/alexmorgan) Tj
0 -25 Td
/F1 13 Tf
(PROFESSIONAL SUMMARY) Tj
/F1 10 Tf
0 -15 Td
(Passionate Full Stack Engineer with experience building scalable web apps with React, Python, and Flask.) Tj
0 -13 Td
(Specializes in REST API architecture, cloud deployments, and responsive UI design.) Tj
0 -25 Td
/F1 13 Tf
(TECHNICAL SKILLS) Tj
/F1 10 Tf
0 -15 Td
(Languages: Python, JavaScript, TypeScript, SQL, HTML5, CSS3) Tj
0 -13 Td
(Frameworks: React, Flask, Node.js, Express, Tailwind CSS) Tj
0 -13 Td
(Databases & Tools: PostgreSQL, MongoDB, Docker, Git, REST APIs, AWS) Tj
0 -25 Td
/F1 13 Tf
(WORK EXPERIENCE) Tj
/F1 11 Tf
0 -15 Td
(Software Engineer Intern - NovaCloud Solutions (Jun 2023 - Present)) Tj
/F1 10 Tf
0 -13 Td
(- Built RESTful microservices in Python Flask, reducing endpoint response times by 28%.) Tj
0 -13 Td
(- Created interactive React dashboards with Tailwind CSS used by 1,000+ daily active users.) Tj
0 -13 Td
(- Built automated CI/CD deployment pipelines using Docker and GitHub Actions.) Tj
0 -25 Td
/F1 13 Tf
(PROJECTS) Tj
/F1 11 Tf
0 -15 Td
(AI Resume Analyzer (React, Flask, Google Gemini API)) Tj
/F1 10 Tf
0 -13 Td
(- Built full-stack app parsing PDF and DOCX documents and analyzing them using Google Gemini AI.) Tj
0 -13 Td
(- Formatted recruiter metrics, keyword scoring, and gap analysis with structured JSON schemas.) Tj
0 -25 Td
/F1 13 Tf
(EDUCATION) Tj
/F1 10 Tf
0 -15 Td
(B.S. in Computer Science - State University of Technology (2020 - 2024, GPA: 3.8/4.0)) Tj
0 -25 Td
/F1 13 Tf
(CERTIFICATIONS) Tj
/F1 10 Tf
0 -15 Td
(- AWS Certified Cloud Practitioner (2024)) Tj
0 -13 Td
(- Meta Certified Front-End Developer (2023)) Tj
ET"""

    stream_bytes = text_content.encode('latin1')
    stream_length = len(stream_bytes)

    objects = [
        # 1: Catalog
        b"<< /Type /Catalog /Pages 2 0 R >>",
        # 2: Pages
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        # 3: Page
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 842] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        # 4: Stream
        f"<< /Length {stream_length} >>\nstream\n".encode('latin1') + stream_bytes + b"\nendstream",
        # 5: Font
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]

    pdf_parts = [b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n"]
    offsets = []

    for i, obj in enumerate(objects, 1):
        offsets.append(sum(len(part) for part in pdf_parts))
        pdf_parts.append(f"{i} 0 obj\n".encode('latin1') + obj + b"\nendobj\n")

    xref_offset = sum(len(part) for part in pdf_parts)
    pdf_parts.append(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode('latin1'))
    for off in offsets:
        pdf_parts.append(f"{off:010d} 00000 n \n".encode('latin1'))

    pdf_parts.append(f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_offset}\n%%EOF\n".encode('latin1'))

    output_path = os.path.join(SAMPLE_DIR, "sample_resume.pdf")
    with open(output_path, "wb") as f:
        f.write(b"".join(pdf_parts))

    print(f"Created sample PDF resume at: {output_path}")
    return output_path


if __name__ == "__main__":
    docx_file = create_sample_docx()
    pdf_file = create_sample_pdf()
    print("Done creating sample resumes!")
